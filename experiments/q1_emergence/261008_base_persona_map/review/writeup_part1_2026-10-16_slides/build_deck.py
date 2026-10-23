"""Build the 16/10 check-in deck (16:9) from the question-based write-up and its title-free figures."""

from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

HERE = Path(__file__).resolve().parent
FIGS = HERE.parent / "writeup_part1_2026-10-16_figs" / "slide"
OUT = HERE / "e3_checkin_2026-10-16.pptx"

FONT = "Charter"
MONO = "Courier New"
INK = RGBColor(0x22, 0x22, 0x22)
TEAL = RGBColor(0x2F, 0x7F, 0x7A)
GREY = RGBColor(0x66, 0x66, 0x66)
BOX = RGBColor(0xF4, 0xF4, 0xF4)
RULE = "BBBBBB"

SW, SH = 13.333, 7.5
LEFT = 0.6
TOP = 1.65  # content starts below the question line and title
EMU_PER_IN = 914400

prs = Presentation()
prs.slide_width = Inches(SW)
prs.slide_height = Inches(SH)
BLANK = prs.slide_layouts[6]


def style_run(run, size, color=INK, bold=False, italic=False, font=FONT):
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color


def textbox(slide, x, y, w, h, anchor=MSO_ANCHOR.TOP):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(0.05)
    tf.margin_top = tf.margin_bottom = Inches(0.03)
    return tb, tf


def set_bullet(p, level=0):
    pPr = p._p.get_or_add_pPr()
    indent = 0.28
    pPr.set("marL", str(Emu(Inches(indent * (level + 1)))))
    pPr.set("indent", str(-Emu(Inches(indent))))
    clr = etree.SubElement(pPr, qn("a:buClr"))
    etree.SubElement(clr, qn("a:srgbClr"), val="2F7F7A")
    etree.SubElement(pPr, qn("a:buFont"), typeface="Arial")
    etree.SubElement(pPr, qn("a:buChar"), char="•" if level == 0 else "–")


def add_paras(tf, items, size, space_after=10, line=1.1):
    """items: str, or (str, dict) with keys level, bullet, color, bold, italic, size, space_before/after."""
    for i, item in enumerate(items):
        text, opt = (item, {}) if isinstance(item, str) else item
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.line_spacing = line
        p.space_after = Pt(opt.get("space_after", space_after))
        if "space_before" in opt:
            p.space_before = Pt(opt["space_before"])
        if opt.get("bullet", True):
            set_bullet(p, opt.get("level", 0))
        r = p.add_run()
        r.text = text
        style_run(
            r,
            opt.get("size", size),
            color=opt.get("color", INK),
            bold=opt.get("bold", False),
            italic=opt.get("italic", False),
        )


def new_slide(title, question=None):
    """Grey question line above a bold title; both sit in one bottom-anchored box."""
    s = prs.slides.add_slide(BLANK)
    _, tf = textbox(s, LEFT, 0.2, SW - 2 * LEFT, 1.35, anchor=MSO_ANCHOR.BOTTOM)
    items = []
    if question:
        items.append((question, {"bullet": False, "color": GREY, "size": 15, "space_after": 4}))
    items.append((title, {"bullet": False, "bold": True, "size": 26, "space_after": 0}))
    add_paras(tf, items, 26, line=1.0)
    _, nf = textbox(s, SW - 1.1, SH - 0.45, 0.6, 0.3)
    nf.paragraphs[0].alignment = PP_ALIGN.RIGHT
    r = nf.paragraphs[0].add_run()
    r.text = str(len(prs.slides._sldIdLst))
    style_run(r, 10, color=GREY)
    return s


def picture(slide, name, x, y, h=None, w=None):
    path = str(FIGS / name)
    if h is not None:
        return slide.shapes.add_picture(path, Inches(x), Inches(y), height=Inches(h))
    return slide.shapes.add_picture(path, Inches(x), Inches(y), width=Inches(w))


def right_of(pic, gap):
    return (pic.left + pic.width) / EMU_PER_IN + gap


def table(slide, rows, x, y, w, col_w, size=14, row_h=0.45):
    gf = slide.shapes.add_table(len(rows), len(rows[0]), Inches(x), Inches(y), Inches(w), Inches(row_h * len(rows)))
    tbl = gf.table
    tblPr = tbl._tbl.tblPr
    tblPr.set("firstRow", "0")
    tblPr.set("bandRow", "0")
    style = tblPr.find(qn("a:tableStyleId"))
    if style is None:
        style = etree.SubElement(tblPr, qn("a:tableStyleId"))
    style.text = "{2D5ABB26-0587-4C30-8999-92F81FD0307C}"  # No Style, No Grid
    for j, cw in enumerate(col_w):
        tbl.columns[j].width = Inches(cw)
    for i, row in enumerate(rows):
        tbl.rows[i].height = Inches(row_h)
        for j, text in enumerate(row):
            cell = tbl.cell(i, j)
            cell.fill.background()
            cell.margin_left = cell.margin_right = Inches(0.08)
            cell.margin_top = cell.margin_bottom = Inches(0.05)
            cell.vertical_anchor = MSO_ANCHOR.TOP
            tf = cell.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            p.line_spacing = 1.0
            r = p.add_run()
            r.text = text
            style_run(r, size, color=TEAL if i == 0 else INK, bold=(i == 0))
            tcPr = cell._tc.get_or_add_tcPr()
            ln = etree.Element(qn("a:lnB"), w="12700", cap="flat", cmpd="sng", algn="ctr")
            sf = etree.SubElement(ln, qn("a:solidFill"))
            etree.SubElement(sf, qn("a:srgbClr"), val="2F7F7A" if i == 0 else RULE)
            etree.SubElement(ln, qn("a:prstDash"), val="solid")
            tcPr.insert(0, ln)
    return tbl


# 1. Title
s = prs.slides.add_slide(BLANK)
_, tf = textbox(s, 1.0, 1.7, SW - 2.0, 2.2, anchor=MSO_ANCHOR.BOTTOM)
add_paras(
    tf,
    [("Can a base language model play a role from a plain-text prompt?", {"bullet": False, "bold": True})],
    40,
    line=1.0,
)
_, tf = textbox(s, 1.0, 4.2, SW - 2.0, 1.2)
add_paras(
    tf,
    [
        ("Check-in, 16 October 2026", {"bullet": False, "color": TEAL, "size": 24, "space_after": 6}),
        ("OLMo-2 7B base, 20 roles", {"bullet": False, "color": GREY, "size": 20}),
    ],
    20,
)

# 2. Summary
s = new_slide("The base model seems to follow example answers much more than a role description")
_, tf = textbox(s, LEFT, TOP + 0.2, SW - 2 * LEFT - 0.5, 3.6)
add_paras(
    tf,
    [
        (
            "When the prompt shows two example answers from a role, the OLMo-2 7B base model writes replies "
            "whose role two blind judges can name about 40% of the time, which passes the bar I set before the run.",
            {"bullet": False},
        ),
        (
            "With only a description of the role it stays under 20%, and when the examples come from a different "
            "role the judges name that other role, so the model seems to follow the examples much more than the "
            "description.",
            {"bullet": False},
        ),
    ],
    22,
    space_after=18,
)
_, tf = textbox(s, LEFT, TOP + 4.1, SW - 2 * LEFT, 0.6)
add_paras(
    tf,
    [("What I'd like from today: answers to the three questions at the end.", {"bullet": False, "color": TEAL})],
    22,
)

# 3. Questions and short answers
s = new_slide("Since 30/09 I worked through nine questions, in this order")
table(
    s,
    [
        ["Date", "Question", "Short answer"],
        [
            "30/09 to 02/10",
            "1. Were my September base-model results valid?",
            "No. The base models got a chat format they had never seen. I withdrew those results.",
        ],
        [
            "05/10",
            "2. Can a small (1B) base model play roles from a plain interview prompt?",
            "Partly. Well above chance, but it fell short on one of three example sets.",
        ],
        [
            "08/10 to 16/10",
            "3. Can the 7B base model play a role from a description alone?",
            "Above chance, but below my bar.",
        ],
        [
            "08/10 to 16/10",
            "4. Does it play the role when the prompt also shows two example answers from that role?",
            "Yes by my rule on the first sample of replies, with a narrow miss on a second sample.",
        ],
        [
            "16/10",
            "5. Is it the examples or the description that makes this work?",
            "Mostly the examples. With another role's examples, the judges name that other role.",
        ],
        [
            "16/10",
            "6. Which roles work?",
            "Roles with a strong voice, like toddler or pirate. Detective and refugee almost never.",
        ],
        [
            "16/10",
            "7. How many replies are broken, and can a judge count them?",
            "A judge marks 5% to 11%, but no judge agreed with my own labels well enough to trust.",
        ],
        ["16/10", "8. Can an automatic judge score role play the way I do?", "Not yet."],
        [
            "16/10",
            "9. Do the model's internal activations carry the role?",
            "With the prompt in view, yes. With the prompt removed, much less. First results only.",
        ],
    ],
    LEFT,
    TOP - 0.05,
    SW - 2 * LEFT,
    [1.75, 5.15, SW - 2 * LEFT - 6.9],
    size=13,
    row_h=0.5,
)

# 4. How I tested the 7B model
s = new_slide(
    "I gave the 7B base model a plain interview transcript in four versions",
    question="How I tested the 7B model (questions 3 to 9)",
)
lw = 6.3
_, tf = textbox(s, LEFT, TOP, lw, 5.4)
add_paras(
    tf,
    [
        "OLMo-2 7B base, plain text only. 20 of the 276 roles, 5 wordings of each description, "
        "18 interview questions: 1,800 replies per prompt.",
        ("Four prompts, each ending with the interview question:", {"space_after": 3}),
        ("the role description only", {"level": 1, "space_after": 1}),
        ("the description plus two example answers from the same role", {"level": 1, "space_after": 1}),
        ("the description plus two example answers from a different role", {"level": 1, "space_after": 1}),
        ("the description plus neutral notes as long as the two examples", {"level": 1}),
        "Two judges, GPT-5.6-Sol and Claude Sonnet 5.5, read only the reply and the 276 role names "
        "and give their top 5 guesses. They never see the prompt.",
        "My bar, written before the run: the lower end of the 95% interval for “right role in the top 5” "
        "is at least 20%; at least 10 of 20 roles beat shuffled role labels; this holds after dropping replies "
        "that copy an example or say the role's name; at most 10% of replies are broken. Both judges must pass.",
    ],
    15,
    space_after=8,
)
rx = LEFT + lw + 0.35
rw = SW - rx - 0.5
_, tf = textbox(s, rx, TOP, rw, 0.35)
add_paras(tf, [("The exact prompt (pirate, two same-role examples, cut short)", {"bullet": False, "color": TEAL})], 13)
prompt_lines = [
    ("<|endoftext|>The following is an interview. The guest answers in their own voice.", False),
    ("About the guest: a pirate captain who has sailed the seven seas in search of treasure and adventure", True),
    ("Interviewer: How should we prepare for economic uncertainties?", True),
    ("Guest: Arrr, matey! When the winds of fortune blow hard and the waves of worry rise, ...", False),
    ("Interviewer: How do you respond when inspiration strikes?", True),
    ("Guest: When inspiration strikes, I embrace it like a long-lost friend. ...", False),
    ("Interviewer: Describe what it's like when you finally understand something complex.", True),
    ("Guest:", False),
]
box, tf = textbox(s, rx, TOP + 0.4, rw, 4.15)
box.fill.solid()
box.fill.fore_color.rgb = BOX
tf.margin_left = tf.margin_right = Inches(0.14)
tf.margin_top = tf.margin_bottom = Inches(0.12)
for i, (line, gap) in enumerate(prompt_lines):
    p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
    p.line_spacing = 1.0
    if gap:
        p.space_before = Pt(9)
    r = p.add_run()
    r.text = line
    style_run(r, 12, font=MONO)
_, tf = textbox(s, rx, TOP + 4.65, rw, 0.5)
add_paras(
    tf,
    [("The description-only prompt is the same without the two example turns.", {"bullet": False, "color": GREY, "italic": True})],
    13,
)

# 5. Questions 3 and 4
s = new_slide(
    "Two same-role examples pass my bar; the description alone does not",
    question="Questions 3 and 4: Can it play a role from the description alone? With two example answers?",
)
pic = picture(s, "fig1_naming_by_prompt.png", 0.45, TOP - 0.05, h=5.75)
rx = right_of(pic, 0.45)
rw = SW - rx - 0.55
_, tf = textbox(s, rx, TOP + 0.3, rw, 5.0)
add_paras(
    tf,
    [
        (
            "Bars show how often the right role is in a judge's top 5 guesses out of 276.",
            {"bullet": False, "color": GREY, "size": 15, "space_after": 16},
        ),
        "Description only: 14.9% and 17.2% (GPT and Claude judges). Well above shuffled labels, "
        "but under my 20% bar.",
        "Two same-role examples: 37.6% and 42.3%, which passes every part of my rule. On a second sample "
        "of replies it narrowly misses: 19.8% for the GPT judge after dropping copied and self-naming replies.",
    ],
    18,
    space_after=16,
)

# 6. Question 5
s = new_slide(
    "With another role's examples, the judges name that other role",
    question="Question 5: Is it the examples or the description that makes this work?",
)
pic = picture(s, "fig2_example_role_wins.png", 0.45, TOP + 0.1, w=8.3)
rx = right_of(pic, 0.4)
rw = SW - rx - 0.5
_, tf = textbox(s, rx, TOP + 0.2, rw, 5.2)
add_paras(
    tf,
    [
        (
            "A chef description with robot examples, asked about good chemistry between people. The reply says chemistry:",
            {"bullet": False, "color": GREY, "size": 15, "space_after": 6},
        ),
        (
            "“can be represented as a binary logic system with three variables”",
            {"bullet": False, "italic": True, "size": 19, "space_after": 6},
        ),
        ("Both judges put robot first.", {"bullet": False, "color": GREY, "size": 15, "space_after": 28}),
        (
            "Length alone does not help. With neutral notes as long as the examples, the right role is in the "
            "top 5 for 8.6% and 12.7% of replies.",
            {"bullet": False},
        ),
    ],
    17,
)

# 7. Question 6
s = new_slide(
    "Roles with a strong voice work; detective and refugee almost never",
    question="Question 6: Which roles work?",
)
pic = picture(s, "fig3_roles_ex2.png", 0.6, TOP - 0.1, h=5.75)
rx = right_of(pic, 0.6)
rw = SW - rx - 0.6
_, tf = textbox(s, rx, TOP + 0.4, rw, 4.8)
add_paras(
    tf,
    [
        "Detective and refugee replies read like ordinary answers. This test can't tell whether the model "
        "plays those roles weakly or the judges just find vivid roles easier to spot.",
        "The assistant is a special case. For 80 of its 90 replies, Claude's first guess is “default”, "
        "which is also one of the 276 names. So a plain assistant-style reply mostly counts as a miss.",
    ],
    18,
    space_after=20,
)

# 8. Question 9
s = new_slide(
    "The activations carry the role with the prompt in view, much less without it",
    question="Question 9: Do the model's internal activations carry the role?",
)
pic = picture(s, "fig6_internal_states_prompt_removed.png", 0.45, TOP - 0.05, h=5.6)
rx = right_of(pic, 0.4)
rw = SW - rx - 0.5
_, tf = textbox(s, rx, TOP, rw, 5.5)
add_paras(
    tf,
    [
        (
            "A simple classifier reads the model's activations and picks which of the 20 roles was prompted. "
            "Chance is 5%.",
            {"bullet": False, "color": GREY, "size": 14, "space_after": 12},
        ),
        "Description only: 52.3% with the prompt in view, 15.2% with the description removed, and 91.8% "
        "from the prompt alone. So I can't yet say the reply itself carries the role.",
        "Same-role examples: 58.6% with the prompt removed, against 44.3% from the reply's words and length. "
        "But the input layer alone already gives 50.1%.",
        (
            "First results. The same-role test on my pre-set rules is still running.",
            {"bullet": False, "color": TEAL, "size": 14, "space_before": 4},
        ),
    ],
    15,
    space_after=12,
)

# 9. Question 7
s = new_slide(
    "A judge marks 5-11% of replies as broken, but I can't trust it yet",
    question="Question 7: How many replies are broken, and can a judge count them?",
)
pic = picture(s, "fig4_broken_replies.png", 0.45, TOP + 0.05, h=5.4)
rx = right_of(pic, 0.4)
rw = SW - rx - 0.5
_, tf = textbox(s, rx, TOP + 0.3, rw, 5.0)
add_paras(
    tf,
    [
        "I wanted a judge's agreement with my own labels (kappa, where 0 is chance and 1 is perfect) to reach "
        "0.6 on both of my label sets. No judge reached it, and these rates are likely too low.",
        "About half the broken replies with same-role examples are the translator translating the question. "
        "One of the translator's two example answers does exactly that.",
    ],
    17,
    space_after=20,
)

# 10. Question 8
s = new_slide(
    "No automatic judge scores role play the way I do yet",
    question="Question 8: Can an automatic judge score role play the way I do?",
)
_, tf = textbox(s, LEFT, TOP + 0.05, SW - 2 * LEFT - 0.5, 1.0)
add_paras(
    tf,
    [
        (
            "I scored 50 replies from 0 to 3 for how well they play a role, before looking at any judge output. "
            "My bar was agreement of at least 0.6 (weighted kappa: 0 is chance, 1 is perfect).",
            {"bullet": False},
        )
    ],
    18,
)
table(
    s,
    [
        ["Comparison", "Agreement [95% interval]", "What it means"],
        ["Jev, a ready-made judge model, vs my first scores", "0.41 [0.21, 0.57]", "Under my 0.6 bar, so I can't use it"],
        [
            "Jev vs my scores after I revised my scale",
            "0.69 [0.49, 0.82]",
            "A first look only, because I had already seen Jev's scores",
        ],
    ],
    LEFT,
    TOP + 1.35,
    SW - 2 * LEFT,
    [4.4, 3.0, SW - 2 * LEFT - 7.4],
    size=18,
    row_h=0.75,
)
_, tf = textbox(s, LEFT, TOP + 4.0, SW - 2 * LEFT - 0.5, 1.0)
add_paras(
    tf,
    [
        (
            "My first scale had a problem too: replies scored against a distant role got 1.57 on average, "
            "more than close roles at 1.13.",
            {"bullet": False, "color": GREY},
        )
    ],
    16,
)

# 11. Questions 1 and 2
s = new_slide(
    "Before this: my September base results were invalid, and the 1B model only partly worked",
    question="Questions 1 and 2: Were my September results valid? Can a small base model play roles?",
)
lw = 4.7
_, tf = textbox(s, LEFT, TOP + 0.1, lw, 5.3)
add_paras(
    tf,
    [
        ("September results", {"bullet": False, "color": TEAL, "bold": True, "size": 16, "space_after": 6}),
        "Blind judges found the right character in their top 5 for 0.7% to 1.4% of base-model replies, "
        "against 1.8% for random guessing. The prompt was a chat format the base models had never seen.",
        "Example turns I added on 14/08 stopped a trained 1B model from playing roles: 55.4% without them, "
        "2.9% with them.",
        ("The 1B base model (right)", {"bullet": False, "color": TEAL, "bold": True, "size": 16, "space_before": 8, "space_after": 6}),
        "Well above the 2% chance level, but it fell short of my rule on example set 2. I treat it as a first look.",
    ],
    15,
    space_after=10,
)
picture(s, "fig5_e0_1b.png", LEFT + lw + 0.35, TOP + 0.05, h=5.45)

# 12. Weak points
s = new_slide(
    "One judge chose the examples and scored the replies, and nobody has checked by hand",
    question="Weak points",
)
_, tf = textbox(s, LEFT, TOP + 0.1, SW - 2 * LEFT - 0.8, 5.3)
add_paras(
    tf,
    [
        "The GPT judge picked the example answers and then scored the replies. Claude agrees with GPT, "
        "but both judges read the same replies.",
        "The 20 roles are a selected subset, and the examples were chosen to be easy to name. "
        "Nothing here is about all 276 roles.",
        "Nobody has checked the role guesses by hand yet.",
        "The same-role result misses my bar on the second sample for one judge by 0.2 points.",
        "The broken-reply definition changed after I saw the data, and no judge passed my agreement bar.",
        "The activation results show the role can be read with the prompt in view. "
        "They don't show where personas come from in training.",
    ],
    19,
    space_after=12,
)

# 13. Plan
s = new_slide(
    "Next week I want to finish the activation results and check the guesses by hand",
    question="Plan for next week (times are rough guesses)",
)
table(
    s,
    [
        ["Task", "Rough time"],
        ["Finish the activation results and write them up", "1 day"],
        ["Check the role guesses by hand on about 100 replies, blind to the judges", "Half a day"],
        [
            "Settle the two open broken-reply cases (toddler replies that tell a story instead of answering, "
            "translator replies that translate the question), then redo the role-guessing scores with the new labels",
            "2 hours",
        ],
        ["Agree the base-model prompt with Ivan, then run the trained models with the same prompt", "2 to 3 days with the GPU machine"],
    ],
    LEFT,
    TOP + 0.1,
    SW - 2 * LEFT,
    [SW - 2 * LEFT - 3.4, 3.4],
    size=18,
    row_h=0.75,
)

# 14. Questions for Ivan
s = new_slide("I'd like your view on three questions before I run the trained models", question="Questions for Ivan")
_, tf = textbox(s, LEFT, TOP + 0.15, SW - 2 * LEFT - 0.4, 5.3)
qs = [
    "The same-role prompt mostly measures how well the base model follows example answers. Is that good enough "
    "for the base model's side of the comparison, or should it use the description alone and accept the lower rate?",
    "If I keep the examples, should the trained models get exactly the same prompt, so only the model changes?",
    "Is the blind role-guessing test enough to say a model plays a role, given that no judge has matched my own "
    "role-play scores yet?",
]
for i, q in enumerate(qs):
    p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
    p.line_spacing = 1.1
    p.space_after = Pt(22)
    pPr = p._p.get_or_add_pPr()
    pPr.set("marL", str(Emu(Inches(0.45))))
    pPr.set("indent", str(-Emu(Inches(0.45))))
    r = p.add_run()
    r.text = f"{i + 1}.\t"
    style_run(r, 21, color=TEAL, bold=True)
    r = p.add_run()
    r.text = q
    style_run(r, 21)

prs.save(OUT)
print(OUT)
